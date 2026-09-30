# -*- coding: utf-8 -*-
"""check_terms.py —— 术语替换三重验证器（solution-doc-standard 技能工具，语言层第 4 步的固化）

用法:
    python -X utf8 check_terms.py <check.json>

check.json 格式:
    {
      "file": "方案.md",              # 单文件；或多文件拼接 "files": ["a.md", "b.md"]
      "body_start": "第一章",          # 正文起点标记：首次出现处之前为文档控制/修订历史区，
                                      #   ABSENT_BODY 只在正文区检查（修订历史行的旧词是合法历史记录）
      "absent_body":  ["治理三角"],    # 旧词必须清零（正文区）
      "present_body": ["三个模块的分工"],  # 新词必须就位（正文区）
      "present_anywhere": ["V1.1"],   # 可选：全文区（含修订历史）检查
      "structure": {"h1": 2, "h2": 4} # 可选：markdown 标题计数不变断言
    }

支持 .md / .txt / .docx（docx 提取段落＋表格文本；structure 断言仅支持 md/txt）。
输出: 逐项 PASS/FAIL 与总结；任何 FAIL 退出码 1。
"""
import argparse
import io
import json
import os
import re
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")


def read_text(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        from docx import Document
        d = Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for r in t.rows:
                for c in r.cells:
                    parts.append(c.text)
        return "\n".join(parts)
    for enc in ("utf-8", "gbk"):
        try:
            return io.open(path, encoding=enc).read()
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise SystemExit(f"无法读取: {path}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("check_json")
    args = ap.parse_args()
    cfg = json.load(io.open(args.check_json, encoding="utf-8"))

    files = cfg.get("files") or ([cfg["file"]] if "file" in cfg else [])
    if not files:
        raise SystemExit("check.json 需要提供 file 或 files")
    full = "\n".join(read_text(f) for f in files)

    results = []  # (ok, label)

    # 正文区切片（排除文档控制/修订历史区）
    body = full
    marker = cfg.get("body_start")
    if marker:
        idx = full.find(marker)
        if idx < 0:
            results.append((False, f"body_start 标记 '{marker}' 未找到（无法定位正文区）"))
        else:
            body = full[idx:]

    for w in cfg.get("absent_body", []):
        n = body.count(w)
        results.append((n == 0, f"ABSENT_BODY  '{w}'：正文区 {n} 处（要求 0）"))

    for w in cfg.get("present_body", []):
        n = body.count(w)
        results.append((n > 0, f"PRESENT_BODY '{w}'：正文区 {n} 处（要求 ≥1）"))

    for w in cfg.get("present_anywhere", []):
        n = full.count(w)
        results.append((n > 0, f"PRESENT_ANY  '{w}'：全文 {n} 处（要求 ≥1）"))

    # 结构不变断言（仅 md/txt：按行首标记计标题数）
    st = cfg.get("structure")
    if st:
        if any(f.lower().endswith(".docx") for f in files):
            results.append((False, "STRUCTURE：docx 不支持标题计数断言，请对源文件做结构检查"))
        else:
            h1 = len(re.findall(r"(?m)^# ", full))
            h2 = len(re.findall(r"(?m)^## ", full))
            if "h1" in st:
                results.append((h1 == st["h1"], f"STRUCTURE   h1：实际 {h1}（要求 {st['h1']}）"))
            if "h2" in st:
                results.append((h2 == st["h2"], f"STRUCTURE   h2：实际 {h2}（要求 {st['h2']}）"))

    fails = 0
    for ok, label in results:
        print(("  PASS " if ok else "  FAIL ") + label)
        fails += (not ok)
    print(f"总结：{len(results) - fails} PASS / {fails} FAIL")
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
